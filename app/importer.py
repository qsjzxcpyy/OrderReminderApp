from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from .date_rules import LOCAL_TZ
from .workflow import suggest_stage


SELECTED_FIELDS = (
    "refrence_no", "platform_user_name", "platform_seller_id", "order_status",
    "date_latest_ship", "tracking_number", "mark_content", "markding_time",
    "system_note", "operator_note", "abnormal_reason", "consignee_name",
    "consignee_country_name", "consignee_email", "warehouse_name",
    "shipping_method_name_cn", "shipping_method_name_en",
)
_DATE_FIELDS = {"date_latest_ship", "markding_time"}


@dataclass
class ImportResult:
    rows: list[dict[str, Any]]
    added: int | None = None
    updated: int | None = None
    duplicates: int = 0
    invalid_rows: list[dict[str, Any]] | None = None
    skipped: int = 0

    def __post_init__(self):
        if self.invalid_rows is None:
            self.invalid_rows = []


def _json_safe(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _parse_datetime(value, field: str):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or LOCAL_TZ).astimezone(LOCAL_TZ)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=LOCAL_TZ)
    text = str(value).strip()
    if text in {"", "0000-00-00 00:00:00"}:
        return None
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, pattern).replace(tzinfo=LOCAL_TZ)
        except ValueError:
            continue
    raise ValueError(f"invalid {field}: {text}")


def _normalize_value(field: str, value):
    if field in _DATE_FIELDS:
        return _parse_datetime(value, field)
    return value if value is None else str(value).strip() if isinstance(value, str) else value


def import_workbook(path: str | Path) -> ImportResult:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("openpyxl is required to import ERP workbooks") from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Export orders" not in workbook.sheetnames:
            raise ValueError("workbook must contain an Export orders worksheet")
        sheet = workbook["Export orders"]
        iterator = iter(sheet.iter_rows(values_only=True))
        try:
            raw_headers = next(iterator)
        except StopIteration:
            raise ValueError("Export orders worksheet is empty")
        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        if "refrence_no" not in headers:
            raise ValueError("Export orders worksheet must contain refrence_no column")

        latest_by_order: dict[str, dict[str, Any]] = {}
        invalid_rows = []
        duplicates = 0
        skipped = 0
        for row_number, values in enumerate(iterator, start=2):
            raw_source = {}
            snapshot_columns = []
            for header, value in zip(headers, values):
                snapshot_columns.append({"header": header, "value": _json_safe(value)})
                if header and header not in raw_source:
                    raw_source[header] = value
            try:
                order_value = values[headers.index("refrence_no")]
                order_no = "" if order_value is None else str(order_value).strip()
                if not order_no:
                    raise ValueError("blank order number")
                normalized = {field: _normalize_value(field, raw_source.get(field)) for field in SELECTED_FIELDS}
                normalized["refrence_no"] = order_no
                normalized["stage_suggestion"] = suggest_stage(normalized)
                normalized["latest_erp_snapshot_json"] = {"columns": snapshot_columns}
            except ValueError as error:
                invalid_rows.append({"row_number": row_number, "reason": str(error)})
                skipped += 1
                continue
            if order_no in latest_by_order:
                duplicates += 1
            latest_by_order[order_no] = normalized
        return ImportResult(
            rows=list(latest_by_order.values()),
            duplicates=duplicates,
            invalid_rows=invalid_rows,
            skipped=skipped,
        )
    finally:
        workbook.close()
